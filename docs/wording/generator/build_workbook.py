#!/usr/bin/env python3
"""Build the Grow Daily wording workbook from the generator's JSON.

Part of the wording inventory (docs/wording). The extractors write their
JSON into generator/.cache (gitignored); this script joins them into one
Excel workbook and then reopens the file to check it.

Inputs, all in the inputs folder (default: generator/.cache):
  dart_rows.json     bin/extract_dart.dart  (--out .cache/dart_rows.json)
  dart_report.json   bin/extract_dart.dart  (--report .cache/dart_report.json)
  native_rows.json   extract_native.py
  usage.json         map_usage.py

Output: docs/wording/GrowDaily-Wording.xlsx (or --out).

Every count in the workbook is a plain value computed here. There are no
formulas, because nothing on this machine can recalculate a workbook
(LibreOffice is not installed), and a formula openpyxl writes has no cached
value until something recalculates it.

One kind of row is added here, beyond the extractors' rows: text members of
class S that return one pattern for both languages, such as
comebackBonusAmount ('+$xp XP'). The Dart extractor skips them because they
have no language branch. Both sides carry the pattern. (The Latin-only rows
with no language branch, such as the "Grow Daily" wordmark or '+$xp XP'
chips outside app_strings.dart, now come from the Dart extractor itself.)

No U+2014 is written anywhere. In wording (Arabic or English) it is the
visible token "[em dash]", so the Flags sheet still shows where the live
copy has one. In reasons (developer comments) it is a comma. After saving,
every text file under docs/wording is checked for the character too.

Usage, from docs/wording/generator:
  python3 build_workbook.py [--inputs DIR] [--out FILE.xlsx]
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import os
import re
import subprocess
import sys
from collections import Counter, OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from map_usage import merge_usage_texts

GEN = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(GEN, '..', '..', '..'))
WORDING_DIR = os.path.abspath(os.path.join(GEN, '..'))
DEFAULT_INPUTS = os.path.join(GEN, '.cache')
DEFAULT_OUT = os.path.join(GEN, '..', 'GrowDaily-Wording.xlsx')

EM_DASH = '\u2014'
EM_TOKEN = '[em dash]'
EXCEL_CELL_MAX = 32767

APP_STRINGS = 'lib/core/l10n/app_strings.dart'
COPY_FILES = ('lib/core/l10n/reminder_copy.dart', 'lib/core/l10n/daily_quotes.dart')

G_APP = 'App strings (app_strings.dart)'
G_REMINDER = 'Notification and reminder copy'
G_QUOTES = 'Daily quotes'
G_INSCREEN = 'In-screen strings (other Dart files)'
G_FAQ = 'Help and Support FAQ'
G_IOS = 'iOS system strings'
G_ANDROID = 'Android system strings'
G_PUSH = 'Push notifications (server)'
G_WEB = 'Web pages'

GROUP_ORDER = [G_APP, G_REMINDER, G_QUOTES, G_INSCREEN, G_FAQ, G_IOS, G_ANDROID, G_PUSH, G_WEB]

# Excel caps sheet names at 31 characters.
GROUP_SHEET = OrderedDict([
    (G_APP, 'App strings (app_strings.dart)'),
    (G_REMINDER, 'Notification and reminder copy'),
    (G_QUOTES, 'Daily quotes'),
    (G_INSCREEN, 'In-screen strings (other Dart)'),
    (G_FAQ, 'Help and Support FAQ'),
    (G_IOS, 'iOS system strings'),
    (G_ANDROID, 'Android system strings'),
    (G_PUSH, 'Push notifications (server)'),
    (G_WEB, 'Web pages'),
])

NO_USAGE_SCREEN = '(no usage found in lib/)'
MANY_SCREENS_PREFIX = 'Many screens, through '
MAX_SCREENS = 10

# Where native, server and web text shows. The Dart rows get their screen
# from map_usage.py (usage.json), which already names every lib file.
NATIVE_SCREENS = {
    'ios/GrowDailyWidget/GrowDailyWidget.swift': 'iOS Home Screen and Lock Screen widgets',
    'ios/GrowDailyWidget/GrowDailyControls.swift': 'iOS Lock Screen and Control Center controls',
    'ios/GrowDailyWidget/GrowDailyAlarmLiveActivity.swift': 'iOS alarm Live Activity',
    'ios/Runner/AlarmKitBridge.swift': 'iOS ringing alarm (AlarmKit)',
    'android/app/src/main/AndroidManifest.xml': 'Android app name (under the icon)',
    'functions/index.js': 'Push notifications: rooms',
    'functions/room_messages.js': 'Push notifications: rooms',
    'public/reset/index.html': 'Web: password reset page (/reset)',
    'public/support.html': 'Web: support page (/support.html)',
    'public/join/index.html': 'Web: room invite page (/join/<code>)',
    'public/open/index.html': 'Web: Lock Screen control fallback page (/open)',
    'public/privacy.html': 'Web: privacy policy (/privacy.html)',
    'web/index.html': 'Web app: browser tab and home screen name',
    'web/manifest.json': 'Web app: install manifest',
}

H_NOTES = 'Generator notes'
COLUMNS = [
    # header, width, wrap, arabic
    ('ID', 34, True, False),
    ('Group', 18, True, False),
    ('Screen / where it shows', 30, True, False),
    ('Section', 22, True, False),
    ('Key', 26, True, False),
    ('Arabic', 48, True, True),
    ('English', 48, True, False),
    ('Reason (why it is worded this way)', 64, True, False),
    ('Defined in (file:line)', 36, True, False),
    ('Used in (file:line)', 42, True, False),
    ('Placeholders', 18, True, False),
    ('Flags', 24, True, False),
    ('Kind', 12, False, False),
    (H_NOTES, 40, True, False),
]

FLAG_ORDER = [
    'em dash',
    'Arabic-Indic digits',
    'Arabic-Indic digits (runtime)',
    'Latin word in Arabic',
    'missing Arabic',
    'missing English',
    'fragment',
    'same text both languages',
    'unused',
    'dynamic',
]

FLAG_MEANING = OrderedDict([
    ('em dash', 'The Arabic or English wording contains an em dash (shown as the token "[em dash]"). '
                'Aziz asked for that dash never to be used in the app\'s copy.'),
    ('Arabic-Indic digits', 'The Arabic wording, as written in the source, contains Arabic-Indic digits '
                            '(U+0660 to U+0669). Informational: listed so each use can be checked, not a '
                            'request to change it.'),
    ('Arabic-Indic digits (runtime)', 'The Arabic wording puts a number through arabicDigits(), so the '
                                      'reader sees Arabic-Indic digits although the source has none: in a '
                                      'placeholder itself ({arabicDigits(n)}), in a local the placeholder '
                                      'names (final d = arabicDigits(done), then {d}), or in a function that '
                                      'does (countedOffsetPhrase, _countedHabits; dart_report.json lists them '
                                      'under digit_functions: the functions that return text and keep the '
                                      'digits, not the ones that turn them back with toWesternDigits or return '
                                      'nothing text can carry). Most of the app shows Western digits, so these '
                                      'are the places that differ.'),
    ('Latin word in Arabic', 'The Arabic wording contains Latin letters. The words found are listed after '
                             'the colon, for example "Latin word in Arabic: Premium". Brand words such as '
                             'Grow Daily, Pro and Premium are listed too, so each can be judged. Placeholders, '
                             'branch conditions ("n == 1:") and, for [dynamic] rows, code are ignored.'),
    ('missing Arabic', 'The Arabic side is empty: the text has no Arabic version (for example English-only '
                       'Swift widget text, a Latin-only literal, or the English-only privacy policy).'),
    ('missing English', 'The English side is empty: the text has no English version (for example the '
                        'Arabic-only room invite page).'),
    ('fragment', 'A piece of a sentence that only one language needs on its own (a local that builds one '
                 'side). It is not flagged missing, because the other side is complete elsewhere.'),
    ('same text both languages', 'Arabic and English are identical. Rows whose only letters are a brand name '
                                 '(Grow Daily, Premium, Pro) or that hold nothing but numbers, symbols and '
                                 'placeholders are not flagged.'),
    ('unused', 'No place in lib/ shows it. For app_strings, reminder_copy and daily_quotes: map_usage.py '
               'found no usage, directly or through another string that is used. For every other Dart row: '
               'following the text with the Dart analyzer reached no widget, notification or other place that '
               'shows it. Either nothing reads the member or field that holds it, or its value only goes to '
               'comparisons, ids, logs or stored data (a row whose text is only written to Firestore or '
               'preferences names those lines in Generator notes). Often left over from a removed feature.'),
    ('dynamic', 'One side is code that could not be written out as text, so the cell shows its source after '
                '"[dynamic] ". Branches, switches, lists and maps are written out instead, one reading per '
                'line.'),
])

BRAND_WORDS = {'grow daily', 'growdaily', 'grow daily premium', 'premium', 'pro'}

LETTER_RE = re.compile(r'[A-Za-z؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]')
ARABIC_INDIC_RE = re.compile(r'[٠-٩]')
LATIN_PHRASE_RE = re.compile(r'[A-Za-z]+(?: [A-Za-z]+)*')
QUOTED_RE = re.compile(r"'((?:[^'\\]|\\.)*)'|\"((?:[^\"\\]|\\.)*)\"")
DART_INTERP_RE = re.compile(r'\$\{[^{}]*\}|\$[A-Za-z_][A-Za-z0-9_]*')
RUNTIME_DIGITS_RE = re.compile(r'\{arabicDigits\(')

FONT = Font(name='Arial', size=10)
FONT_BOLD = Font(name='Arial', size=10, bold=True)
FONT_HEADER = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FONT_TITLE = Font(name='Arial', size=14, bold=True)
FONT_SECTION = Font(name='Arial', size=11, bold=True)
FILL_HEADER = PatternFill('solid', start_color='2E5E4E', end_color='2E5E4E')
FILL_SECTION = PatternFill('solid', start_color='E3EFE9', end_color='E3EFE9')
BORDER_HEADER = Border(bottom=Side(style='thin', color='1B3A30'))

ALIGN_LTR_WRAP = Alignment(horizontal='left', vertical='top', wrap_text=True)
ALIGN_LTR = Alignment(horizontal='left', vertical='top', wrap_text=False)
ALIGN_RTL_WRAP = Alignment(horizontal='right', vertical='top', wrap_text=True, readingOrder=2)
ALIGN_CODE_WRAP = Alignment(horizontal='left', vertical='top', wrap_text=True, readingOrder=1)
ALIGN_HEADER = Alignment(horizontal='left', vertical='center', wrap_text=True)
ALIGN_HEADER_RTL = Alignment(horizontal='right', vertical='center', wrap_text=True, readingOrder=2)
ALIGN_NUM = Alignment(horizontal='right', vertical='top')

LINE_PT = 12.75
MAX_ROW_LINES = 12
MAX_USED_IN = 15


# ─── Text helpers ────────────────────────────────────────────────────────────

def wording_text(s: str) -> str:
    return (s or '').replace(EM_DASH, EM_TOKEN)


def comment_text(s: str) -> str:
    """Developer comments keep their words; an em dash becomes a comma."""
    if not s:
        return ''
    out = []
    for line in s.split('\n'):
        line = re.sub(r'^[ \t]*\u2014[ \t]*', '', line)
        line = re.sub(r'[ \t]*\u2014[ \t]*', ', ', line)
        out.append(line)
    return '\n'.join(out)


def cell_text(s):
    if not isinstance(s, str):
        return s
    s = ILLEGAL_CHARACTERS_RE.sub('', s)
    if EM_DASH in s:
        raise SystemExit('build_workbook: an em dash reached a cell: ' + s[:120])
    if len(s) > EXCEL_CELL_MAX:
        note = ' [cut at the Excel cell limit]'
        s = s[:EXCEL_CELL_MAX - len(note)] + note
    return s


def strip_placeholders(s: str) -> str:
    prev = None
    while prev != s:
        prev = s
        s = re.sub(r'\{[^{}]*\}', '', s)
    return s


def visible_text(side: str, visible: str | None) -> str:
    """The part of a side a reader can see: the literal text only, without
    branch conditions, placeholders or, for [dynamic] code, the code."""
    if visible is not None:
        s = visible
    else:
        s = side or ''
        if s.startswith('[dynamic]'):
            code = s[len('[dynamic]'):]
            parts = []
            for m in QUOTED_RE.finditer(code):
                after = code[m.end():m.end() + 4].lstrip()
                before = code[max(0, m.start() - 4):m.start()].rstrip()
                if after.startswith('=>') or after.startswith(':') or before.endswith('=='):
                    continue  # a switch pattern, a map key or a comparison, not text
                parts.append(m.group(1) if m.group(1) is not None else m.group(2))
            s = '\n'.join(parts)
            s = DART_INTERP_RE.sub('', s)
    s = s.replace(EM_TOKEN, ' ')
    return strip_placeholders(s)


def compute_flags(row: dict, member_unused: bool) -> list[str]:
    ar, en = row['arabic'], row['english']
    ar_vis = visible_text(ar, row.get('arabic_visible'))
    en_vis = visible_text(en, row.get('english_visible'))
    flags = []
    if EM_TOKEN in ar or EM_TOKEN in en or EM_DASH in ar or EM_DASH in en:
        flags.append('em dash')
    if ARABIC_INDIC_RE.search(ar):
        flags.append('Arabic-Indic digits')
    if ar.strip() and (RUNTIME_DIGITS_RE.search(ar) or row.get('arabic_digits_runtime')):
        flags.append('Arabic-Indic digits (runtime)')
    latin = []
    for m in LATIN_PHRASE_RE.findall(ar_vis):
        if m not in latin:
            latin.append(m)
    if latin:
        flags.append('Latin word in Arabic: ' + ', '.join(latin))
    if row.get('fragment'):
        flags.append('fragment')
    else:
        if not ar.strip():
            flags.append('missing Arabic')
        if not en.strip():
            flags.append('missing English')
    if ar.strip() and en.strip():
        na, ne = ' '.join(ar.split()), ' '.join(en.split())
        if na == ne:
            letters = ' '.join(''.join(ch if LETTER_RE.match(ch) else ' ' for ch in ar_vis).split())
            if letters and letters.lower() not in BRAND_WORDS:
                flags.append('same text both languages')
    if member_unused:
        flags.append('unused')
    if row.get('dynamic') or ar.startswith('[dynamic]') or en.startswith('[dynamic]'):
        flags.append('dynamic')
    return flags


def flag_name(flag: str) -> str:
    return flag.split(':')[0]


# ─── Inputs ──────────────────────────────────────────────────────────────────

def load_json(path: str):
    if not os.path.exists(path):
        raise SystemExit(f'build_workbook: missing input {path}. Run the extractors first (see README.md).')
    with open(path, encoding='utf-8') as fh:
        return json.load(fh)


def read_lines(rel: str) -> list[str]:
    with open(os.path.join(REPO, rel), encoding='utf-8') as fh:
        return fh.read().split('\n')


def comment_block_above(lines: list[str], line_no: int) -> str:
    """Contiguous // and /// lines right above line_no, joined like the extractor does."""
    i = line_no - 2
    block = []
    while i >= 0 and lines[i].strip().startswith('//'):
        block.append(lines[i].strip())
        i -= 1
    block.reverse()
    paras, cur = [], []
    for ln in block:
        text = re.sub(r'^///?', '', ln).strip()
        if not text:
            if cur:
                paras.append(' '.join(cur))
                cur = []
            continue
        cur.append(text)
    if cur:
        paras.append(' '.join(cur))
    return '\n'.join(paras)


def dart_pattern(literal: str) -> tuple[str, list[str]]:
    names = []

    def repl(m):
        tok = m.group(0)
        name = tok[2:-1] if tok.startswith('${') else tok[1:]
        names.append(name)
        return '{' + name + '}'

    return DART_INTERP_RE.sub(repl, literal), names


def extra_rows(dart_rows, usage) -> tuple[list[dict], list[str]]:
    """S members that return one pattern for both languages; see the module
    docstring."""
    rows, notes = [], []
    covered = {r['id'].split('#')[0].split('[')[0] for r in dart_rows}
    lines = read_lines(APP_STRINGS)
    added = 0
    skipped_members = []
    for name, m in usage['members'].items():
        if f'app_strings.{name}' in covered or name in ('_cutoffClockEn', '_cutoffClockAr'):
            continue
        end = m['line'] - 1
        while end < len(lines) - 1 and not lines[end].rstrip().endswith((';', '}')):
            end += 1
        src = '\n'.join(lines[m['line'] - 1:end + 1])
        if 'isAr' in src or 'isArabic' in src:
            notes.append(f'S.{name} has a language test but no row; left out, check the Dart extractor')
            continue
        # Only a body that is one string literal and nothing else, such as
        # `=> '$tier · $family';`. A helper like _byAmount, which picks an
        # Arabic particle with a RegExp, is grammar inside other strings.
        body = re.search(r"=>\s*('(?:[^'\\]|\\.)*'|\"(?:[^\"\\]|\\.)*\")\s*;\s*$", src)
        if not body:
            skipped_members.append(name)
            continue
        raw = body.group(1)[1:-1]
        text, names = dart_pattern(raw)
        text_line = m['line'] + src[:body.start(1)].count('\n')
        rows.append({
            'id': f'app_strings.{name}',
            'source_group': G_APP,
            'section': '(one pattern for both languages)',
            'key': name,
            'arabic': text,
            'english': text,
            'placeholders': names,
            'reason': comment_block_above(lines, m['line']),
            'file': APP_STRINGS,
            'line': m['line'],
            'text_line': text_line,
            'ar_line': text_line,
            'en_line': text_line,
            'kind': m['kind'],
            'notes': ['One pattern for both languages: the member has no language branch.'],
            '_added': 'language-neutral S member',
        })
        added += 1
    notes.append(f'{added} language-neutral S members added (one pattern for both languages)')
    if skipped_members:
        notes.append('S members with no row and no single literal body, left out: '
                     + ', '.join(skipped_members))
    return rows, notes


# ─── Where a string shows ────────────────────────────────────────────────────

class Where:
    """Screens and usage sites for a row, from usage.json."""

    def __init__(self, usage):
        self.members = usage['members']
        self.top = usage['top_level']
        self.rows = usage.get('dart_rows', {})
        self.file_to_screen = usage['file_to_screen']

    @staticmethod
    def _caller(internal: str) -> tuple[str, str]:
        m = re.match(r'^(.*?):(\d+) \((.+)\)$', internal)
        return (m.group(1), m.group(3)) if m else ('', '')

    def _resolve(self, table, key, stem_of_caller, chain=()):
        e = table.get(key)
        if not e:
            return [], []
        if e['usages']:
            via = ' > '.join(chain)
            uses = [u + (f' (via {via})' if via else '') for u in e['usages']]
            return list(e['screens']), uses
        screens, uses = [], []
        for internal in e.get('internal_usages', []):
            file, caller = self._caller(internal)
            if not caller:
                continue
            ckey = stem_of_caller(file, caller)
            if ckey == key or ckey in [stem_of_caller(file, c) for c in chain]:
                continue
            s, u = self._resolve(table, ckey, stem_of_caller, chain + (caller,))
            screens += [x for x in s if x not in screens]
            uses += [x for x in u if x not in uses]
        return screens, uses

    def notes_for(self, row) -> list[str]:
        """Remarks map_usage.py made about how the row's usage was found."""
        entry = self.rows.get(row['id'])
        return list(entry.get('notes') or []) if entry else []

    def for_row(self, row):
        """Returns (screens, used_in, member_unused, mapped).

        mapped: the row's usage was looked up, so no screen means no usage."""
        f = row['file']
        if f == APP_STRINGS:
            base = row['key'].split('[')[0]
            names = [base] if base in self.members else [n for n in (base + 'En', base + 'Ar') if n in self.members]
            screens, uses, unused = [], [], bool(names)
            for n in names:
                s, u = self._resolve(self.members, n, lambda _f, c: c)
                screens += [x for x in s if x not in screens]
                uses += [x for x in u if x not in uses]
                unused = unused and bool(self.members[n].get('unused'))
            return screens, uses, unused, True
        if f in COPY_FILES:
            stem = os.path.basename(f)[:-5]
            key = f"{stem}.{row['key'].split('[')[0]}"
            if key in self.top:
                s, u = self._resolve(self.top, key, lambda cf, c: f'{os.path.basename(cf)[:-5]}.{c}')
                return s, u, bool(self.top[key].get('unused')), True
        entry = self.rows.get(row['id'])
        if row['source_group'] == G_ANDROID and f.endswith('.dart'):
            return (['Android system Settings: Apps, Grow Daily, Notifications'],
                    list(entry['usages']) if entry else [], False, False)
        if f == 'ios/Runner/Info.plist':
            if row['section'] == 'App name':
                return ['iOS app name (under the icon)'], [], False, False
            if entry and entry['screens']:
                return (['iOS permission prompt, over: ' + '; '.join(entry['screens'])],
                        entry['usages'], False, False)
            return ['iOS permission prompts'], [], False, False
        if entry is not None:
            if entry['usages']:
                return list(entry['screens']), list(entry['usages']), False, False
            if entry.get('unused'):
                return [NO_USAGE_SCREEN], [], True, False
        if f in self.file_to_screen:
            return [self.file_to_screen[f]], list(row.get('used_in') or []), False, False
        if f in NATIVE_SCREENS:
            return [NATIVE_SCREENS[f]], list(row.get('used_in') or []), False, False
        return [f], list(row.get('used_in') or []), False, False


def defined_in(r: dict) -> str:
    """"file:line" of the declaration, then where each side's own text is
    when that is somewhere else."""
    base = f"{r['file']}:{r['line']}"
    sides = []
    for side, has in (('ar', r['arabic'].strip()), ('en', r['english'].strip())):
        line = r.get(f'{side}_line')
        if not has or not line:
            continue
        path = r.get(f'{side}_file') or r['file']
        sides.append((side, path, line))
    if not sides or all(p == r['file'] and ln == r['line'] for _, p, ln in sides):
        return base
    return base + '\n' + '; '.join(f'{side}: {p}:{ln}' for side, p, ln in sides)


# ─── Git ─────────────────────────────────────────────────────────────────────

def git_info() -> dict:
    def run(*args):
        try:
            return subprocess.run(['git', '-C', REPO, *args], capture_output=True, text=True,
                                  check=True).stdout
        except Exception:
            return ''
    head = run('rev-parse', '--short', 'HEAD').strip() or 'unknown'
    status = [ln for ln in run('status', '--porcelain').split('\n') if ln.strip()]
    tracked = [ln for ln in status if not ln.startswith('??')]
    untracked = [ln for ln in status if ln.startswith('??')]
    under_app = [ln for ln in tracked if ln[3:].startswith(('lib/', 'ios/', 'android/', 'functions/',
                                                              'public/', 'web/'))]
    return {
        'head': head,
        'dirty': bool(status),
        'tracked_changes': len(tracked),
        'untracked': len(untracked),
        'app_changes': [ln[3:] for ln in under_app],
    }


# ─── Workbook writing ────────────────────────────────────────────────────────

def estimate_lines(text: str, width: int) -> int:
    if not text:
        return 1
    per_line = max(8, int(width * 1.15))
    return sum(max(1, math.ceil(len(p) / per_line)) for p in str(text).split('\n'))


def write_data_sheet(ws, records):
    for ci, (header, width, _wrap, arabic) in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=ci, value=header)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.border = BORDER_HEADER
        c.alignment = ALIGN_HEADER_RTL if arabic else ALIGN_HEADER
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 30
    for ri, rec in enumerate(records, 2):
        lines = 1
        for ci, (header, width, wrap, arabic) in enumerate(COLUMNS, 1):
            value = cell_text(rec[header])
            c = ws.cell(row=ri, column=ci, value=value)
            if isinstance(value, str) and value.startswith('='):
                c.data_type = 's'
            c.font = FONT
            if arabic:
                # Source code reads left to right; RTL reordering scrambles it.
                code = isinstance(value, str) and value.startswith('[dynamic]')
                c.alignment = ALIGN_CODE_WRAP if code else ALIGN_RTL_WRAP
            else:
                c.alignment = ALIGN_LTR_WRAP if wrap else ALIGN_LTR
            if wrap and value:
                lines = max(lines, estimate_lines(value, width))
        ws.row_dimensions[ri].height = round(min(lines, MAX_ROW_LINES) * LINE_PT + 3, 2)
    ws.freeze_panes = 'A2'
    last = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f'A1:{last}{max(1, len(records) + 1)}'
    ws.sheet_view.zoomScale = 110


def set_cell(ws, row, col, value, font=FONT, align=ALIGN_LTR_WRAP, fill=None):
    c = ws.cell(row=row, column=col, value=cell_text(value))
    if isinstance(c.value, str) and c.value.startswith('='):
        c.data_type = 's'
    c.font = font
    c.alignment = align
    if fill is not None:
        c.fill = fill
    return c


def header_row(ws, row, values, rtl_cols=()):
    for ci, v in enumerate(values, 1):
        c = set_cell(ws, row, ci, v, FONT_HEADER, ALIGN_HEADER_RTL if ci in rtl_cols else ALIGN_HEADER,
                     FILL_HEADER)
        c.border = BORDER_HEADER


def write_summary(ws, records, group_counts):
    ws.column_dimensions['A'].width = 46
    for col in 'BCDEFGHIJKLM':
        ws.column_dimensions[col].width = 16
    r = 1
    set_cell(ws, r, 1, 'Summary', FONT_TITLE, ALIGN_LTR)
    r += 1
    set_cell(ws, r, 1, 'Every number here is a value computed when the workbook was generated. '
                       'None is a formula, so nothing needs recalculating.', FONT, ALIGN_LTR)
    r += 2

    # Rows per source group.
    set_cell(ws, r, 1, 'Rows per source group', FONT_SECTION, ALIGN_LTR)
    r += 1
    set_cell(ws, r, 1, '"With a reason" counts rows whose Reason column holds a developer comment. '
                       'Remarks the generator adds (where a comment came from, how two sides were '
                       'paired) are in the Generator notes column and are not counted.', FONT, ALIGN_LTR)
    r += 1
    header_row(ws, r, ['Source group', 'Rows', 'With Arabic', 'With English', 'With a reason',
                       'With generator notes', 'With a usage (Used in)', 'Flagged rows', 'Unused'])
    ws.row_dimensions[r].height = 30
    r += 1
    totals = [0] * 8
    for g in GROUP_ORDER:
        recs = [x for x in records if x['Group'] == g]
        vals = [len(recs),
                sum(1 for x in recs if x['Arabic'].strip()),
                sum(1 for x in recs if x['English'].strip()),
                sum(1 for x in recs if x['Reason (why it is worded this way)'].strip()),
                sum(1 for x in recs if x[H_NOTES].strip()),
                sum(1 for x in recs if x['Used in (file:line)'].strip()),
                sum(1 for x in recs if x['Flags']),
                sum(1 for x in recs if 'unused' in x['_flag_names'])]
        set_cell(ws, r, 1, g)
        for i, v in enumerate(vals):
            set_cell(ws, r, 2 + i, v, FONT, ALIGN_NUM)
            totals[i] += v
        r += 1
    set_cell(ws, r, 1, 'Total', FONT_BOLD)
    for i, v in enumerate(totals):
        set_cell(ws, r, 2 + i, v, FONT_BOLD, ALIGN_NUM)
    r += 3

    # Top screens. "Many screens, through X" is not a screen: it stands for
    # a member that reaches more than 12 screens, so it is counted apart.
    screen_counts = Counter()
    screen_groups = {}
    many = Counter()
    no_usage = 0
    for x in records:
        if x['_screens'] == [NO_USAGE_SCREEN]:
            no_usage += 1
            continue
        for s in x['_screens']:
            if s.startswith(MANY_SCREENS_PREFIX):
                many[s[len(MANY_SCREENS_PREFIX):]] += 1
                continue
            screen_counts[s] += 1
            screen_groups.setdefault(s, Counter())[x['Group']] += 1
    set_cell(ws, r, 1, 'Top 40 screens by number of strings', FONT_SECTION, ALIGN_LTR)
    r += 1
    many_text = ''
    if many:
        many_text = (' Not counted as a screen either: "Many screens, through X", which stands for a member '
                     'that reaches more than 12 screens (' + '; '.join(
                         f'{k}: {n} strings' for k, n in sorted(many.items(), key=lambda kv: (-kv[1], kv[0])))
                     + ').')
    set_cell(ws, r, 1, 'A string shown on several screens counts once on each. '
                       f'Not counted: {no_usage} strings with no usage found in lib/.{many_text} '
                       f'Distinct screens or places in total: {len(screen_counts)}.', FONT, ALIGN_LTR)
    r += 1
    header_row(ws, r, ['Screen / where it shows', 'Strings', 'Rank', 'Groups'])
    r += 1
    ranked = sorted(screen_counts.items(), key=lambda kv: (-kv[1], kv[0].lower()))[:40]
    for rank, (s, n) in enumerate(ranked, 1):
        set_cell(ws, r, 1, s)
        set_cell(ws, r, 2, n, FONT, ALIGN_NUM)
        set_cell(ws, r, 3, rank, FONT, ALIGN_NUM)
        groups = '; '.join(f'{g} ({c})' for g, c in screen_groups[s].most_common())
        set_cell(ws, r, 4, groups, FONT, ALIGN_LTR)
        r += 1
    r += 2

    # Flags.
    flag_counts = Counter(f for x in records for f in x['_flag_names'])
    set_cell(ws, r, 1, 'Rows per flag', FONT_SECTION, ALIGN_LTR)
    r += 1
    set_cell(ws, r, 1, 'A row can carry several flags. '
                       f'Rows with at least one flag: {sum(1 for x in records if x["Flags"])}.', FONT, ALIGN_LTR)
    r += 1
    header_row(ws, r, ['Flag', 'Rows'] + [GROUP_SHEET[g] for g in GROUP_ORDER])
    ws.row_dimensions[r].height = 42
    r += 1
    for f in FLAG_ORDER:
        set_cell(ws, r, 1, f)
        set_cell(ws, r, 2, flag_counts.get(f, 0), FONT, ALIGN_NUM)
        for gi, g in enumerate(GROUP_ORDER):
            n = sum(1 for x in records if x['Group'] == g and f in x['_flag_names'])
            set_cell(ws, r, 3 + gi, n, FONT, ALIGN_NUM)
        r += 1
    r += 2
    set_cell(ws, r, 1, 'What each flag means', FONT_SECTION, ALIGN_LTR)
    r += 1
    header_row(ws, r, ['Flag', 'Meaning'])
    r += 1
    for f, meaning in FLAG_MEANING.items():
        set_cell(ws, r, 1, f)
        set_cell(ws, r, 2, meaning, FONT, ALIGN_LTR)
        r += 1
    ws.freeze_panes = None


def write_readme(ws, ctx):
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 120
    r = 1
    set_cell(ws, r, 1, 'Grow Daily wording inventory', FONT_TITLE, ALIGN_LTR)
    ws.row_dimensions[r].height = 22
    r += 2

    def section(title):
        nonlocal r
        set_cell(ws, r, 1, title, FONT_SECTION, ALIGN_LTR, FILL_SECTION)
        set_cell(ws, r, 2, '', FONT, ALIGN_LTR, FILL_SECTION)
        r += 1

    def item(label, text):
        nonlocal r
        set_cell(ws, r, 1, label, FONT_BOLD, ALIGN_LTR_WRAP)
        set_cell(ws, r, 2, text, FONT, ALIGN_LTR_WRAP)
        ws.row_dimensions[r].height = round(max(estimate_lines(label, 30), estimate_lines(text, 120))
                                            * LINE_PT + 3, 2)
        r += 1

    section('What this file is')
    item('Purpose', 'Every piece of wording a person can read in Grow Daily, in Arabic and English, with the '
                    'reason the developers wrote down for it and where it lives. One row per wording.')
    item('Covers', 'The Flutter app (lib/): class S in app_strings.dart, notification and reminder copy, daily '
                   'quotes, the Help and Support FAQ, strings written directly in screens, models, catalogs and '
                   'services, Latin-only text with no language branch, and the Android notification channel '
                   'names. Outside Dart: iOS Info.plist and Swift (widgets, controls, the alarm Live Activity), '
                   'the Android manifest, room push notifications from functions/, and the web pages in '
                   'public/ (the privacy policy line by line) and web/.')
    item('Reason', 'Copied from the source: the doc comments (///) and // comments directly above the '
                   'declaration first, then the comments on the lines that hold the text, plus a trailing '
                   'comment on the same line, a comment written inside the language test itself (between isAr '
                   'and ?, or between ? and the text), and the comment above the statement, switch arm, '
                   'widget or call that holds the text. Empty when the developers wrote nothing. It is the developers\' '
                   'own explanation, not an interpretation: anything the generator has to say (where a '
                   'comment came from, how two sides were paired, why a text has no Arabic) is in the '
                   'Generator notes column instead. Decorative box-drawing rules around a comment heading are '
                   'dropped.')
    r += 1

    section('Generation')
    item('Generated', ctx['generated'])
    item('Git commit', f"{ctx['git']['head']}")
    dirty = ctx['git']
    if dirty['dirty']:
        app = dirty['app_changes']
        tree = (f"Dirty: {dirty['tracked_changes']} tracked files modified and {dirty['untracked']} untracked "
                f"paths. The extractors read the working tree, so uncommitted edits ARE included.")
        if app:
            shown = ', '.join(app[:20]) + (f' and {len(app) - 20} more' if len(app) > 20 else '')
            tree += f' Modified app sources: {shown}.'
    else:
        tree = 'Clean: the workbook matches the commit above.'
    item('Working tree', tree)
    item('Inputs', ctx['inputs_text'])
    item('Rows', ctx['rows_text'])
    item('Numbers are values', 'Every count in this workbook (Summary sheet, this sheet) is a plain value '
                               'computed when the file was generated. There are no formulas: LibreOffice is '
                               'not installed on the machine that builds it, so nothing could recalculate '
                               'them, and a formula written by openpyxl has no value until something does. '
                               'Regenerate the file to refresh the numbers.')
    r += 1

    section('How to regenerate')
    item('Commands', ctx['commands'])
    item('One line', ctx['one_line'])
    item('What each step does',
         '1. dart pub get --offline: resolves the analyzer package from the local pub cache.\n'
         '2. bin/extract_dart.dart: parses every .dart file under lib/ (parse only, nothing is built) and '
         'writes the Dart rows plus a report of what it skipped.\n'
         '3. extract_native.py: reads ios/, android/, functions/, public/ and web/ and writes the native rows.\n'
         '4. map_usage.py: finds where each S member and each reminder_copy and daily_quotes declaration is '
         'used, and names the screen of every lib file. It also resolves lib/ with the Dart analyzer '
         '(bin/verify_usage.dart), which records where every referenced value and every string literal goes '
         'next (returned, stored in a field, passed on, shown, written to storage, or dropped), and follows '
         'every other Dart row\'s text along that to the screens that draw it. With --verify it also checks '
         'its own S usage map against the analyzer.\n'
         '5. build_workbook.py: joins the JSON into this workbook, then reopens it and checks it.\n'
         'Intermediate JSON goes to docs/wording/generator/.cache, which git ignores. Nothing under lib/, '
         'test/, ios/, android/, functions/, public/ or web/ is written.')
    r += 1

    section('Sheets')
    item('Read me', 'This sheet.')
    item('Summary', 'Rows per source group, the 40 screens with the most strings, and rows per flag. '
                    '"Many screens, through X" is not a screen, so it is left out of the ranking and of the '
                    'number of distinct screens and named apart.')
    item('All wording', 'Every row, sorted by group, then rows with a usage before rows with none, then '
                        'screen, file and line.')
    for g in GROUP_ORDER:
        item(GROUP_SHEET[g], f'Only the rows of the group "{g}", same columns and order.')
    item('Flags', 'Only the rows that carry at least one flag, same columns and order.')
    r += 1

    section('Columns')
    for label, text in [
        ('ID', 'Stable id: source file name without extension, a dot, then the key. "#n" is added when one key '
               'gives several rows (for example several branches of one method), and "[case]" names the '
               'case of a switch or the element of a list (PrayerCalcMethod.label[jordan]), including a '
               'language ?: inside one arm of a switch (HabitCue._presetLabel[fajr]). Ids hold no '
               'spaces, parentheses or operators: a condition in a case key is spelled out (h < 12 becomes '
               'h_lt_12). Native ids start with ios., android., push. or web. Ids ending in ".latin_only" '
               'are Latin text with no language branch.'),
        ('Group', 'Where the wording comes from. Groups are listed in this order: ' + ', '.join(GROUP_ORDER) + '.'),
        ('Screen / where it shows', 'For app_strings, reminder_copy and daily_quotes: the screens of every file '
                                    'that uses the string, from map_usage.py; a string used only by another '
                                    'string gets the screens of that one. For every other Dart row: the places '
                                    'the text itself reaches, found by following it with the Dart analyzer. A '
                                    'catalog entry (IslamicHabitTemplate, AchievementModel, CharacterOption) is '
                                    'followed from the fields its constructor arguments fill, not from the list '
                                    'that holds the entry, so .length or .id on the list counts for nothing. A '
                                    'value is followed when it is returned (to the members that call that one), '
                                    'stored in a field (to that field\'s reads), or passed to a parameter (to '
                                    'that parameter\'s reads, and back to the same call when the callee hands '
                                    'it back). It shows where it is handed to a widget, a notification, the home '
                                    'widget or a thrown exception. It stops, with no screen, where it is written '
                                    'to Firestore or preferences, or used for something that is not text (a '
                                    'comparison, a condition, a length, a log line). So an action or a write, '
                                    'such as DashboardNotifier._resolveUnlocks or RoomsController._profileFields, '
                                    'is never a screen of its own: only the widgets that draw the value are (the '
                                    '"Warrior" fallback name shows on the Room page through '
                                    'RoomParticipant.displayName). A method that returns nothing text can carry '
                                    '(void, bool), or whose text is shown by the method itself (a thrown '
                                    'exception, a notification), shows it where it is called. A member read on '
                                    'one enum constant (MilestoneType.levelUp.localizedName) counts only for '
                                    'that constant\'s row, and so does one read inside a switch case that has '
                                    'already narrowed the receiver (case MilestoneType.levelUp: then '
                                    'e.type.localizedName(isAr)). A read of only a list\'s size '
                                    '(kDailyQuotes.length) is not a use of its text. Text in a build method with '
                                    'no caller shows where it '
                                    'is written, and the screen is that file\'s. A Riverpod provider, or a '
                                    'member that reaches more than 25 places, is named once: "Many screens, '
                                    'through X" past 12 screens. '
                                    f'"{NO_USAGE_SCREEN}" means no place that shows it was found. '
                                    'iOS permission prompts name the screens whose action opens them. For '
                                    'other native, server and web rows: the surface that shows it, named from '
                                    f'the file. At most {MAX_SCREENS} screens are listed, then "and N more".'),
        ('Section', 'The nearest "// ── X ──" banner above the text in the source (a banner that wraps over '
                    'several comment lines counts), or the class name. For native rows, the part of the file '
                    '(for example "Permission prompts (NS*UsageDescription)").'),
        ('Key', 'Getter, method, field or constant name, Class.member outside app_strings.dart, or the plist, JSON '
                'or HTML key.'),
        ('Arabic', 'The Arabic text exactly as written. Interpolations appear as {name}, or as the code that '
                   'computes the value when it is not a plain name ({daysCount(counted)}, '
                   '{quoted.join(\'، \')}). When a side reads '
                   'differently under different conditions, each reading is on its own line after its '
                   'condition, in source order: "n == 1: ...", "n <= 10: ...", "otherwise: ..." (a switch '
                   'case, an if, a ?: or a ?? fallback); a reading that is an empty string shows as "(empty)". '
                   'A local that builds a piece of the sentence is written '
                   'under it as "name =" with its own readings, including a local that is itself a language '
                   'test (done = {arabicDigits(weekDone)} on the Arabic side, done = {weekDone} on the English). '
                   'A language test with no letter on either side (only placeholders, digits and punctuation) is '
                   'no row when both sides are the same or when it is only a piece of a member whose other rows '
                   'are wording; a member whose whole result is one keeps its row. When the two sides line up case by case (an '
                   'enum\'s names, a list), each case is its own row instead. A side that is code which cannot '
                   'be written out shows its source after "[dynamic] " and is left aligned. An em dash in the '
                   'live copy appears as "[em dash]".'),
        ('English', 'Same rules as Arabic.'),
        ('Reason (why it is worded this way)', 'The developers\' comments, as written: the declaration\'s doc '
                                               'comment first, then the comments that sit with the text: above '
                                               'the lines that hold it, inside the language test (between isAr '
                                               'and ?, or between ? and the text), and above the statement, '
                                               'switch arm, widget or call that holds it (a comment above an if '
                                               'counts for every row inside it; a statement with no comment of '
                                               'its own takes the one that heads its run of statements, when no '
                                               'blank line separates them). A comment of a local written under a '
                                               'sentence goes with that sentence. Paragraphs are kept as line breaks. An em dash in a '
                                               'comment is written as a comma, because this workbook never '
                                               'contains that character.'),
        ('Defined in (file:line)', 'First line: the repo-relative path and the 1-based line where the '
                                   'declaration that holds the wording starts (after its doc comment and '
                                   'annotations; for Swift, the func, var or struct around the text). When '
                                   'the text itself sits on other lines, or in another file, a second line '
                                   'gives each side\'s place: "ar: path:line; en: path:line", the line of the '
                                   'sentence the cell starts with, not of a local written under it (for example the '
                                   'reset page\'s Arabic HTML and its English script dictionary, Info.plist '
                                   'and ar.lproj/InfoPlist.strings, or a Swift func and the line of its '
                                   'text).'),
        ('Used in (file:line)', 'Each place that uses the text, one line per file:line, at most ' + str(MAX_USED_IN)
                                + ' then "+N more". "(via X)" or "(via X > Y)" means it is reached through X: a '
                                'string or member, a field the text is stored in (RoomParticipant.displayName) '
                                'or a parameter it is passed to (toWesternDigits(input)). When several chains '
                                'reach the same line they share it, shortest first: "(via X; or via Y > Z)", or '
                                '"(directly; or via X)" when the text also reaches it with no step between. "(in X, which reaches '
                                'N places)" marks a member named once instead of place by place. For push '
                                'rows, the lines of functions/index.js that send '
                                'the message. Empty when the text is written in place (a build method, a web '
                                'page, a widget), where "Defined in" is where it shows.'),
        ('Placeholders', 'The interpolated names, comma separated.'),
        ('Flags', 'Zero or more flags, separated by "; ". See below.'),
        ('Kind', 'getter, method, field, constructor-args, map-entry, plist, xml, js, html or other.'),
        (H_NOTES, 'Remarks from the generator, not from the developers: where a Reason taken from somewhere '
                  'other than the line above came from, how the English and Arabic of a web page were paired, '
                  'why a Latin-only text has no Arabic, which lines use an Android channel name, which fields '
                  'a catalog entry\'s screens were read from, and the write lines of text that only goes to '
                  'stored data.'),
    ]:
        item(label, text)
    r += 1

    section('Flags')
    for f, meaning in FLAG_MEANING.items():
        item(f, meaning)
    r += 1

    section('Notes and known gaps')
    for label, text in ctx['notes']:
        item(label, text)


# ─── Verification ────────────────────────────────────────────────────────────

def verify(out_path, records, group_counts, flagged_count) -> bool:
    print('\nVerification: reopening', out_path)
    wb = load_workbook(out_path)
    ok = True
    expected = {'All wording': len(records), 'Flags': flagged_count}
    for g in GROUP_ORDER:
        expected[GROUP_SHEET[g]] = group_counts.get(g, 0)
    print('Sheets and data rows:')
    for ws in wb.worksheets:
        rows = ws.max_row - 1 if ws.title in expected else None
        tag = ''
        if rows is not None:
            if rows != expected[ws.title]:
                ok = False
                tag = f'  MISMATCH, expected {expected[ws.title]}'
            tag += f'  (autofilter {ws.auto_filter.ref}, frozen at {ws.freeze_panes})'
            if not ws.auto_filter.ref or ws.freeze_panes != 'A2':
                ok = False
            headers = [c.value for c in ws[1]]
            if headers != [c[0] for c in COLUMNS]:
                ok = False
                tag += '  HEADERS DIFFER'
        print(f'  {ws.title}: {rows if rows is not None else ws.max_row} {"data rows" if rows is not None else "rows used"}{tag}')

    ws = wb['All wording']
    headers = [c.value for c in ws[1]]
    ix = {h: i for i, h in enumerate(headers)}
    by_id = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        by_id.setdefault(row[ix['ID']], []).append(row)
    missing, dup, empty_ar, empty_en, changed = [], [], [], [], []
    for rec in records:
        got = by_id.get(rec['ID'])
        if not got:
            missing.append(rec['ID'])
            continue
        if len(got) > 1:
            dup.append(rec['ID'])
        row = got[0]
        if rec['Arabic'] and not row[ix['Arabic']]:
            empty_ar.append(rec['ID'])
        if rec['English'] and not row[ix['English']]:
            empty_en.append(rec['ID'])
        if (row[ix['Arabic']] or '') != cell_text(rec['Arabic']) or (row[ix['English']] or '') != cell_text(rec['English']):
            changed.append(rec['ID'])
    print(f'All wording: {len(by_id)} distinct ids, missing {len(missing)}, duplicated {len(dup)}')
    print(f'Source has Arabic but the cell is empty: {len(empty_ar)}; source has English but the cell is empty: {len(empty_en)}')
    print(f'Cells whose text differs from the source (after em dash tokenising): {len(changed)}')
    if missing or dup or empty_ar or empty_en or changed:
        ok = False
        print('  examples:', (missing + dup + empty_ar + empty_en + changed)[:10])
    bad_ids = [i for i in by_id if not re.fullmatch(r'[A-Za-z0-9_.#\[\]:\-]+', str(i))]
    print(f'Ids with a space, parenthesis or other odd character: {len(bad_ids)}')
    ok = ok and not bad_ids
    repeated = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        places = [ln.split(' ')[0] for ln in (row[ix['Used in (file:line)']] or '').split('\n')
                  if ln and not ln.startswith('+')]
        if len(places) != len(set(places)):
            repeated.append(row[ix['ID']])
    print(f'Used in cells that name the same file:line twice: {len(repeated)} {repeated[:3]}')
    ok = ok and not repeated

    em, formulas = 0, 0
    for w in wb.worksheets:
        for row in w.iter_rows():
            for c in row:
                if isinstance(c.value, str) and EM_DASH in c.value:
                    em += 1
                if c.data_type == 'f':
                    formulas += 1
    print(f'Cells with an em dash: {em}; formula cells: {formulas}')
    ok = ok and em == 0 and formulas == 0

    ar_col = ix['Arabic'] + 1
    bad_align = 0
    for ri in range(2, ws.max_row + 1):
        c = ws.cell(row=ri, column=ar_col)
        code = isinstance(c.value, str) and c.value.startswith('[dynamic]')
        want = ('left', 1) if code else ('right', 2)
        if (c.alignment.horizontal, c.alignment.readingOrder) != want:
            bad_align += 1
    fnt = ws.cell(row=2, column=1).font
    print(f'Arabic cells with the wrong alignment: {bad_align}; body font {fnt.name} {fnt.sz}')
    ok = ok and bad_align == 0 and fnt.name == 'Arial'

    # The em dash rule covers the whole deliverable folder, not only the workbook.
    folder_hits = []
    for root, dirs, files in os.walk(WORDING_DIR):
        dirs[:] = [d for d in dirs if d != '.dart_tool']
        for name in files:
            p = os.path.join(root, name)
            if name.endswith('.xlsx'):
                continue
            try:
                with open(p, encoding='utf-8') as fh:
                    if EM_DASH in fh.read():
                        folder_hits.append(os.path.relpath(p, WORDING_DIR))
            except (UnicodeDecodeError, OSError):
                continue
    print(f'Text files under docs/wording with an em dash: {len(folder_hits)} {folder_hits[:5]}')
    ok = ok and not folder_hits

    print('\nFive sample rows:')
    wanted_groups = [G_APP, G_REMINDER, G_INSCREEN, G_IOS, G_PUSH]
    for g in wanted_groups:
        cand = [r for r in ws.iter_rows(min_row=2, values_only=True)
                if r[ix['Group']] == g and r[ix['Reason (why it is worded this way)']]
                and r[ix['Arabic']] and len(r[ix['English']] or '') > 15]
        if not cand:
            cand = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[ix['Group']] == g]
        if not cand:
            continue
        row = cand[len(cand) // 3]

        def short(v, n=110):
            v = '' if v is None else str(v).replace('\n', ' | ')
            return v if len(v) <= n else v[:n] + '...'
        print(f"  [{row[ix['ID']]}]")
        for h in ('Group', 'Screen / where it shows', 'Section', 'Arabic', 'English',
                  'Reason (why it is worded this way)', 'Defined in (file:line)', 'Used in (file:line)',
                  'Placeholders', 'Flags', 'Kind', H_NOTES):
            print(f'    {h}: {short(row[ix[h]])}')
    print('\nVerification', 'PASSED' if ok else 'FAILED')
    return ok


# ─── Main ────────────────────────────────────────────────────────────────────

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--inputs', default=DEFAULT_INPUTS,
                    help='folder holding dart_rows.json, dart_report.json, native_rows.json and usage.json')
    ap.add_argument('--out', default=DEFAULT_OUT, help='workbook to write')
    args = ap.parse_args()
    inputs = os.path.abspath(args.inputs)
    out_path = os.path.abspath(args.out)

    paths = {n: os.path.join(inputs, n) for n in
             ('dart_rows.json', 'dart_report.json', 'native_rows.json', 'usage.json')}
    dart_rows = load_json(paths['dart_rows.json'])
    report = load_json(paths['dart_report.json'])
    native_rows = load_json(paths['native_rows.json'])
    usage = load_json(paths['usage.json'])

    # Consistency: the usage map must describe the same app_strings.dart as the rows.
    stale = [r['id'] for r in dart_rows if r['file'] == APP_STRINGS
             and r['key'] in usage['members'] and usage['members'][r['key']]['line'] != r.get('decl_line', r['line'])]
    if stale:
        raise SystemExit(f'build_workbook: usage.json and dart_rows.json disagree on {len(stale)} line numbers '
                         f'(for example {stale[:3]}). Rerun the extractors so they read the same sources.')
    rows_meta = usage.get('_meta', {}).get('dart_rows_usage', {})
    missing_usage = [r['id'] for r in dart_rows if r['file'] != APP_STRINGS and r['file'] not in COPY_FILES
                     and r['id'] not in usage.get('dart_rows', {})
                     and (r.get('member') or {}).get('kind') in ('getter', 'method', 'function', 'field')]
    if rows_meta.get('rows_mapped') is not None and missing_usage:
        raise SystemExit(f'build_workbook: usage.json has no entry for {len(missing_usage)} Dart rows '
                         f'(for example {missing_usage[:3]}). Rerun map_usage.py after extract_dart.dart.')

    added_rows, added_notes = extra_rows(dart_rows, usage)
    all_rows = dart_rows + added_rows + native_rows
    ids = Counter(r['id'] for r in all_rows)
    dupes = [i for i, n in ids.items() if n > 1]
    if dupes:
        raise SystemExit(f'build_workbook: duplicate ids {dupes[:10]}')
    unknown = {r['source_group'] for r in all_rows} - set(GROUP_ORDER)
    if unknown:
        raise SystemExit(f'build_workbook: unknown source groups {unknown}')

    where = Where(usage)
    records = []
    for r in all_rows:
        row = dict(r)
        row['arabic'] = wording_text(r['arabic'])
        row['english'] = wording_text(r['english'])
        screens, uses, member_unused, mapped = where.for_row(r)
        # One line per file:line, whichever chains reach it.
        uses = merge_usage_texts(uses)
        if mapped and not screens:
            screens = [NO_USAGE_SCREEN]
        flags = compute_flags(row, member_unused)
        used_in = uses[:MAX_USED_IN]
        if len(uses) > MAX_USED_IN:
            used_in.append(f'+{len(uses) - MAX_USED_IN} more')
        shown_screens = screens[:MAX_SCREENS]
        screen_text = '; '.join(shown_screens)
        if len(screens) > MAX_SCREENS:
            screen_text += f'; and {len(screens) - MAX_SCREENS} more'
        records.append({
            'ID': r['id'],
            'Group': r['source_group'],
            'Screen / where it shows': screen_text,
            'Section': comment_text(r.get('section', '')),
            'Key': r['key'],
            'Arabic': row['arabic'],
            'English': row['english'],
            'Reason (why it is worded this way)': comment_text(r.get('reason', '')),
            'Defined in (file:line)': defined_in(r),
            'Used in (file:line)': '\n'.join(used_in),
            'Placeholders': ', '.join(r.get('placeholders') or []),
            'Flags': '; '.join(flags),
            'Kind': r.get('kind', ''),
            H_NOTES: comment_text('\n'.join(list(r.get('notes') or []) + where.notes_for(r))),
            '_screens': screens,
            '_flag_names': [flag_name(f) for f in flags],
            '_sort': (GROUP_ORDER.index(r['source_group']),
                      1 if screens == [NO_USAGE_SCREEN] else 0,
                      '; '.join(screens).lower(), r['file'], int(r['line']),
                      int(r.get('text_line') or r['line']), r['id']),
        })
    records.sort(key=lambda x: x['_sort'])
    group_counts = Counter(x['Group'] for x in records)
    flagged = [x for x in records if x['Flags']]

    git = git_info()
    now = dt.datetime.now().astimezone()

    def input_line(name):
        p = paths[name]
        m = dt.datetime.fromtimestamp(os.path.getmtime(p)).strftime('%Y-%m-%d %H:%M')
        return f'{name} (written {m})'
    analyzer = usage.get('_meta', {}).get('analyzer_check', {})
    if analyzer.get('ran'):
        agree = (analyzer.get('s_usages_analyzer') == analyzer.get('s_usages_python')
                 and analyzer.get('no_external_usage_sets_equal')
                 and analyzer.get('top_level_usages_analyzer') == analyzer.get('top_level_usages_python'))
        analyzer_text = ('The S usage map was cross-checked with the Dart analyzer (map_usage.py --verify): '
                         + ('both agree' if agree else 'they DISAGREE, see usage.json _meta.analyzer_check')
                         + f", {analyzer.get('s_usages_python')} usages of S members and "
                         f"{analyzer.get('top_level_usages_python')} of the top-level copy.")
    else:
        analyzer_text = 'The S usage map was not cross-checked with the Dart analyzer (run map_usage.py --verify).'
    if rows_meta.get('rows_mapped') is not None:
        analyzer_text += (f" The analyzer mapped {rows_meta['rows_mapped']} other Dart and native rows: "
                          f"{rows_meta['rows_with_usages']} with usages, {rows_meta['rows_in_place']} written "
                          f"in place, {len(rows_meta['rows_no_reference_found'])} with no place that shows them "
                          f"({len(rows_meta.get('rows_only_stored', []))} of those only written to stored "
                          "data).")
    else:
        analyzer_text += (' The other Dart rows were NOT mapped with the analyzer ('
                          + str(rows_meta.get('why', 'map_usage.py --no-analyzer'))[:200]
                          + '), so they only name their own file\'s screen.')

    gen_dir = GEN
    commands = (f'cd {gen_dir}\n'
                'dart pub get --offline\n'
                'dart run bin/extract_dart.dart --out .cache/dart_rows.json --report .cache/dart_report.json\n'
                'python3 extract_native.py\n'
                'python3 map_usage.py --verify\n'
                'python3 build_workbook.py\n'
                'Timings on the machine that built this file: extract_dart.dart about 8 s, extract_native.py '
                'under 1 s, map_usage.py about 50 s and build_workbook.py about 4 s. map_usage.py runs the '
                'Dart analyzer once whether or not --verify is given, and --verify reuses that run, so it '
                'adds almost nothing. --no-analyzer skips the analyzer (about 19 s), but then the rows outside '
                'app_strings, reminder_copy and daily_quotes only name their own file\'s screen and --verify '
                'is ignored.')
    one_line = (f'cd {gen_dir} && dart pub get --offline && dart run bin/extract_dart.dart --out '
                '.cache/dart_rows.json --report .cache/dart_report.json && python3 extract_native.py && '
                'python3 map_usage.py --verify && python3 build_workbook.py')

    summary_counts = '; '.join(f'{g}: {group_counts.get(g, 0)}' for g in GROUP_ORDER)
    report_arabic = len(report.get('uncovered_arabic_literals', []))
    report_skipped = len(report.get('skipped_conditionals', []))
    latin_rows = sum(1 for r in dart_rows if r.get('latin_only'))
    notes = [
        ('Rows added by this step',
         'Some S members return one pattern for both languages (for example comebackBonusAmount, "+{xp} XP"); '
         'they are rows here with the pattern on both sides and Section "(one pattern for both languages)". '
         '_byAmount is not a row: it only picks the Arabic particle for other strings. Counts: '
         + '; '.join(added_notes) + '.'),
        ('Latin-only rows', f'{latin_rows} rows (ids ending in ".latin_only") are Latin text with no language '
                            'branch, so an Arabic reader sees it as written: the "Grow Daily" wordmark, "PRO" and '
                            '"Premium", the "+{xp} XP" chips, "LVL {level}", AM and PM in the plan picker, the '
                            'font names, the medal numerals I to IV, the "App:", "Platform:" and "Language:" lines '
                            'of the support mail body, the "Warrior" display name written to the account and to '
                            'the room leaderboard (app_strings.dart says the app no longer invents that name, but '
                            'AuthNotifier.initialDisplayName and RoomsController._profileFields still do), the '
                            '"Grow Daily" stand-in habit name of a snoozed reminder, and exception messages that '
                            'may reach the user. Their Arabic side is empty, so they carry "missing Arabic".'),
        ('Not in any row', f'{report_arabic} Arabic literals in lib/ are not wording and have no row: keyword '
                           'matchers (habit category and step detection, dhikr detection), the moderation '
                           'blocklist, cue synonyms and the Arabic-Indic digit table. The list is in '
                           'dart_report.json under uncovered_arabic_literals. Language conditionals that hold no '
                           'wording have no row either. All of them are in dart_report.json under '
                           f'skipped_conditionals ({report_skipped} entries), each with its reason: no string '
                           'literal in either branch (the text comes from elsewhere), join pieces (a comma that joins a '
                           'list, isAr ? \'، \' : \', \', and the «و» / "and" that HabitCue._timeLabel joins its '
                           'last time with), pieces with no letter on either side (isAr ? \'"$n"\' : \'"$n"\', '
                           'the \'$today$tasks\' body of streakRiskCopy, the arabicDigits(minutes) fallback of '
                           'reminderOffsetLabel), locale code pairs, the "d MMMM" / "MMM d" date format pattern '
                           'and an internal theme id named "preview". The «و» / "And " that _oneQuitSentence '
                           'puts before a second sentence is no row either, but it is written under both of its '
                           'sentences as "and =". extract_native.py skips server error '
                           'messages the app never shows, log text and ids; its coverage lists are in '
                           'native_summary.json.'),
        ('Branches and lists', 'A side with several readings lists each after its condition, in source order. '
                               'A local that holds a piece of a sentence, including a local that is a language '
                               'test of its own (habitOnTimeLine\'s target, done and needed), is written under '
                               'every sentence that uses it, on the side that reads it. '
                               'The conditions are the source\'s own ("n == 0 || (n >= 3 && n <= 10)"), so the '
                               'plural rules the reasons explain can be checked line by line. A conditional '
                               'inside an interpolation is written out the same way, one reading per branch: '
                               '"n == 1: {n} sprint logged today" / "otherwise: {n} sprints logged today", or '
                               '"appVersion == null: App: Grow Daily" / "otherwise: App: Grow Daily '
                               '{appVersion}". When the Arabic has more forms than the English (1, 2, 3 to 10, '
                               '11 and up against one and many), the Arabic forms that the same English form '
                               'covers share one row.'),
        ('Swift is English only', 'Every Swift string (widgets, controls, the Live Activity) has no Arabic: the '
                                  'widget extension has no .lproj folder. Those rows are flagged "missing Arabic".'),
        ('Android', 'There is no strings.xml. The native Android wording is the app label in AndroidManifest.xml. '
                    'The notification channel names and descriptions that Android lists under Settings, Apps, '
                    'Grow Daily, Notifications ("Grow Daily", "Habit reminders and progress celebrations", '
                    '"Grow Daily alarms", "Reminders you chose to ring as alarms") are defined in Dart '
                    '(notification_service.dart) and are rows in the Android system strings group. They are '
                    'English only.'),
        ('Usage map limits', analyzer_text + ' Only lib/ is scanned, so a string used only by tests counts as '
                             'unused. A file that holds several surfaces (for example profile_screen_sheets.dart) '
                             'has one screen name. Doc-comment references such as [S.tagline] are not usages. '
                             'A member called only through a supertype or a dynamic call has no resolved '
                             'reference and would show as unused. Text is not followed through stored data: '
                             'a value written to Firestore or preferences and read back by key somewhere else is '
                             'not traced to that reader, so text that only goes to storage is flagged unused (the '
                             'catalog habit descriptions: IslamicHabitTemplate.localDescription has no caller, '
                             'and description is only copied into other templates, HabitModel and Firestore). '
                             'A value handed to a callback held in a variable (onPick(name)) or to a Riverpod '
                             'notifier\'s state counts as shown where it is handed over. A parameter is followed '
                             'the same way for every call, except that a callee which hands the parameter '
                             'straight back continues at the call that passed it.'),
        ('Row heights', f'Rows are sized for their text up to {MAX_ROW_LINES} lines so the sheets stay easy to '
                        'scan. Select a cell to read a longer reason in full, or drag the row border.'),
        ('Section oddities', 'Section is the nearest banner, which inside app_strings.dart is sometimes only '
                             'loosely related (for example _cutoffClock sits under "Tasbih").'),
        ('Possible issues seen while extracting',
         'The brand is spelled two ways. "GrowDaily" (no space): the English iOS permission sheets in '
         'Info.plist, the privacy and support pages, and the web app shell and manifest. "Grow Daily": '
         'CFBundleDisplayName and CFBundleName, ar.lproj/InfoPlist.strings, the Android app label, the reset, '
         'join and open pages, and the app itself. '
         'The room push stand-ins "Someone" and "your room" (functions/index.js) are English and are spliced '
         'into the Arabic pushes of events A and C and the nudge, and "your room" into the Arabic habit-added '
         'push too. '
         'Two join separators always use the Arabic comma whatever the language '
         '(room_detail_screen_header_progress.dart and room_detail_screen_participant_calendar.dart). The '
         'widget gallery preview data includes "mohdabood2003", which looks like a real user handle.'),
    ]

    ctx = {
        'generated': now.strftime('%Y-%m-%d %H:%M %Z').strip(),
        'git': git,
        'inputs_text': (f'From {inputs}: ' + ', '.join(input_line(n) for n in paths) + '.'),
        'rows_text': (f'{len(records)} rows in total ({summary_counts}). {len(flagged)} rows carry at least one '
                      'flag. See the Summary sheet for more.'),
        'commands': commands,
        'one_line': one_line,
        'notes': notes,
    }

    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = 'Read me'
    write_readme(ws_readme, ctx)
    write_summary(wb.create_sheet('Summary'), records, group_counts)
    write_data_sheet(wb.create_sheet('All wording'), records)
    for g in GROUP_ORDER:
        write_data_sheet(wb.create_sheet(GROUP_SHEET[g]), [x for x in records if x['Group'] == g])
    write_data_sheet(wb.create_sheet('Flags'), flagged)
    wb.properties.title = 'Grow Daily wording inventory'
    wb.properties.creator = 'docs/wording/generator/build_workbook.py'

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    print(f'wrote {out_path}')
    print(f'rows: {len(records)} ({summary_counts}); flagged: {len(flagged)}')
    fc = Counter(f for x in records for f in x['_flag_names'])
    print('flags: ' + '; '.join(f'{f}: {fc.get(f, 0)}' for f in FLAG_ORDER))
    for n in added_notes:
        print('note:', n)

    summary = OrderedDict([
        ('rows', len(records)),
        ('by_group', OrderedDict((g, group_counts.get(g, 0)) for g in GROUP_ORDER)),
        ('flagged', len(flagged)),
        ('flags', OrderedDict((f, fc.get(f, 0)) for f in FLAG_ORDER)),
        ('with_reason', sum(1 for x in records if x['Reason (why it is worded this way)'].strip())),
        ('with_used_in', sum(1 for x in records if x['Used in (file:line)'].strip())),
    ])
    with open(os.path.join(inputs, 'workbook_summary.json'), 'w', encoding='utf-8') as fh:
        json.dump(summary, fh, ensure_ascii=False, indent=1)

    if not verify(out_path, records, group_counts, len(flagged)):
        sys.exit(1)


if __name__ == '__main__':
    main()
